#!/usr/bin/env python3
"""md2cv — turn structured Markdown résumés into neatly typeset DOCX files.

md2cv targets one specific résumé shape; it is NOT a general Markdown→DOCX
converter. Anything outside the expected patterns becomes a plain paragraph:
readable, but not specially typeset. If you need a universal converter, use
pandoc instead.

Expected Markdown shape:
    # Full Name                       -> name
    tagline / contact lines           -> header block (until first "## ")
    ## Section Title                  -> section heading + rule
    ### Company                       -> company
    **Role** · dates                  -> role line (dates right-aligned)
    #### Project: X                   -> sub-project (bold italic)
    - bullet                          -> list item

Inline emphasis: **bold**, *italic* and _italic_ (underscores are not
recognised mid-word, so snake_case_word stays literal).

Unsupported but tolerated (as plain text):
    numbered lists, tables, fenced code blocks, block quotes, links,
    images, headings deeper than ####, YAML frontmatter, and any role/date
    separator other than "·" (U+00B7 MIDDLE DOT).
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass, fields
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

__version__ = "0.3.0"
PROG = "md2cv"


# --------------------------------------------------------------- patterns ----

ROLE_RE = re.compile(r"^\*\*(.+?)\*\*\s*·\s*(.+)$")
INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|(?<!\w)_[^_]+_(?!\w))")


# ----------------------------------------------------------------- presets ---

PRESETS = {
    "default": {},
    "compact": {
        "font": "Arial Narrow",
        "body_size": 9.2,
        "line_spacing": 1.0,
        "margin_top": 1.0,
        "margin_bottom": 1.0,
        "margin_left": 1.3,
        "margin_right": 1.3,
        "section_space_before": 7.0,
        "company_space_before": 6.0,
        "item_space_after": 0.8,
    },
    "comfortable": {
        "body_size": 10.0,
        "line_spacing": 1.15,
        "margin_top": 1.5,
        "margin_bottom": 1.5,
        "margin_left": 2.0,
        "margin_right": 2.0,
        "section_space_before": 12.0,
        "company_space_before": 10.0,
        "item_space_after": 2.0,
        "para_space_after": 3.0,
    },
    "classic": {
        "accent": "#2F5496",
        "rule_size": 0,
    },
}


# ----------------------------------------------------------------- config ----


@dataclass
class Config:
    # typography
    font: str = "Calibri"
    body_size: float = 9.5
    name_size: float = 20.0
    tagline_size: float = 10.5
    contact_size: float = 9.0
    section_size: float = 11.0
    company_size: float = 11.5
    role_size: float = 10.0
    project_size: float = 10.5
    line_spacing: float = 1.0
    # page geometry (cm)
    page_width: float = 21.0
    page_height: float = 29.7
    margin_top: float = 1.1
    margin_bottom: float = 1.1
    margin_left: float = 1.5
    margin_right: float = 1.5
    # colors
    accent: str = "#1F3864"
    muted: str = "#595959"
    rule_color: str = "#9AA7B4"
    rule_size: int = 6
    # spacing (pt / cm)
    section_space_before: float = 9.0
    section_space_after: float = 2.5
    company_space_before: float = 8.0
    item_space_after: float = 1.0
    para_space_after: float = 1.5
    list_indent: float = 0.5

    @property
    def text_width(self) -> float:
        return self.page_width - self.margin_left - self.margin_right


# ---------------------------------------------------------------- helpers ----


def _rgb(hex_str: str) -> RGBColor:
    s = hex_str.lstrip("#")
    return RGBColor(int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16))


def _hex(value: str) -> str:
    s = value.lstrip("#")
    if len(s) != 6 or any(c not in "0123456789abcdefABCDEF" for c in s):
        raise argparse.ArgumentTypeError(
            f"expected a hex color like #RRGGBB, got {value!r}"
        )
    return "#" + s.upper()


def _set_run(run, size, font, *, bold=False, italic=False, color=None):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(attr), font)
    return run


def _add_bottom_rule(paragraph, *, sz: int, color: str, space: int = 3):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(sz))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color.lstrip("#"))
    pBdr.append(bottom)
    successors = (
        "w:shd",
        "w:tabs",
        "w:suppressAutoHyphens",
        "w:kinsoku",
        "w:wordWrap",
        "w:overflowPunct",
        "w:topLinePunct",
        "w:autoSpaceDE",
        "w:autoSpaceDN",
        "w:bidi",
        "w:adjustRightInd",
        "w:snapToGrid",
        "w:spacing",
        "w:ind",
        "w:contextualSpacing",
        "w:mirrorIndents",
        "w:suppressOverlap",
        "w:jc",
        "w:textDirection",
        "w:textAlignment",
        "w:textboxTightWrap",
        "w:outlineLvl",
        "w:divId",
        "w:cnfStyle",
        "w:rPr",
        "w:sectPr",
        "w:pPrChange",
    )
    pPr.insert_element_before(pBdr, *successors)


# --------------------------------------------------------------- renderer ----


class Renderer:
    def __init__(self, cfg: Config):
        self.cfg = cfg
        self.accent = _rgb(cfg.accent)
        self.muted = _rgb(cfg.muted)
        self._name: str | None = None
        self.doc = self._setup()

    def _setup(self) -> Document:
        c = self.cfg
        doc = Document()
        normal = doc.styles["Normal"]
        normal.font.name = c.font
        normal.font.size = Pt(c.body_size)
        rf = normal.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(attr), c.font)
        pf = normal.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = c.line_spacing

        for s in doc.sections:
            s.page_width = Cm(c.page_width)
            s.page_height = Cm(c.page_height)
            s.top_margin = Cm(c.margin_top)
            s.bottom_margin = Cm(c.margin_bottom)
            s.left_margin = Cm(c.margin_left)
            s.right_margin = Cm(c.margin_right)
            s.header_distance = Cm(0.6)
            s.footer_distance = Cm(0.6)
        return doc

    def _inline(self, p, text, *, size, bold=False, italic=False, color=None):
        for part in INLINE_RE.split(text):
            if not part:
                continue
            b, i = bold, italic
            if part.startswith("**") and part.endswith("**") and len(part) > 4:
                part, b = part[2:-2], True
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                part, i = part[1:-1], True
            elif part.startswith("_") and part.endswith("_") and len(part) > 2:
                part, i = part[1:-1], True
            _set_run(
                p.add_run(part), size, self.cfg.font, bold=b, italic=i, color=color
            )

    def name(self, text):
        self._name = text
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        self._inline(p, text, size=self.cfg.name_size, bold=True)

    def tagline(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        self._inline(p, text, size=self.cfg.tagline_size)

    def contact(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0.5)
        self._inline(p, text, size=self.cfg.contact_size, color=self.muted)

    def section(self, title):
        c = self.cfg
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(c.section_space_before)
        pf.space_after = Pt(c.section_space_after)
        pf.keep_with_next = True
        _set_run(
            p.add_run(title.upper()),
            c.section_size,
            c.font,
            bold=True,
            color=self.accent,
        )
        if c.rule_size > 0:
            _add_bottom_rule(p, sz=c.rule_size, color=c.rule_color)

    def company(self, name):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(self.cfg.company_space_before)
        pf.space_after = Pt(0.5)
        pf.keep_with_next = True
        self._inline(p, name, size=self.cfg.company_size, bold=True)

    def role_line(self, role, dates):
        c = self.cfg
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(2)
        pf.keep_with_next = True
        pf.tab_stops.add_tab_stop(Cm(c.text_width), WD_TAB_ALIGNMENT.RIGHT)
        _set_run(p.add_run(role), c.role_size, c.font, bold=True)
        _set_run(p.add_run("\t" + dates), c.contact_size, c.font, color=self.muted)

    def project(self, title):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(5)
        pf.space_after = Pt(1)
        pf.keep_with_next = True
        self._inline(p, title, size=self.cfg.project_size, bold=True, italic=True)

    def bullet(self, text):
        c = self.cfg
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(c.list_indent)
        pf.first_line_indent = Cm(-c.list_indent)
        pf.space_after = Pt(c.item_space_after)
        pf.tab_stops.add_tab_stop(Cm(c.list_indent), WD_TAB_ALIGNMENT.LEFT)
        _set_run(p.add_run("•\t"), c.body_size, c.font)
        self._inline(p, text, size=c.body_size)

    def paragraph(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(self.cfg.para_space_after)
        self._inline(p, text, size=self.cfg.body_size)

    def render(self, text: str) -> Document:
        lines = [l.rstrip() for l in text.splitlines()]
        n = len(lines)
        i, in_header, expect_role = 0, True, False

        while i < n:
            s = lines[i].strip()
            if not s:
                i += 1
                continue

            # Role line is expected right after ### Company or #### Project,
            # possibly separated by blank lines. Any other non-empty line
            # cancels the expectation.
            if expect_role:
                m = ROLE_RE.match(s)
                expect_role = False
                if m:
                    self.role_line(m.group(1), m.group(2))
                    i += 1
                    continue

            if s.startswith("# ") and not s.startswith("## "):
                self.name(s[2:])
                i += 1
                continue
            if s.startswith("## "):
                in_header = False
                self.section(s[3:])
                i += 1
                continue
            if s.startswith("#### "):
                self.project(s[5:])
                expect_role = True
                i += 1
                continue
            if s.startswith("### "):
                self.company(s[4:])
                expect_role = True
                i += 1
                continue
            if in_header:
                (self.tagline if s.startswith("**") else self.contact)(s)
                i += 1
                continue
            if s.startswith("- "):
                self.bullet(s[2:])
                i += 1
                continue
            self.paragraph(s)
            i += 1

        if self._name:
            self.doc.core_properties.author = self._name
            self.doc.core_properties.title = self._name
        return self.doc


# -------------------------------------------------------------------- io -----


def convert(src: Path, dst: Path, cfg: Config) -> None:
    text = src.read_text(encoding="utf-8-sig")  # tolerate BOM
    Renderer(cfg).render(text).save(str(dst))


def _resolve_output(src: Path, args) -> Path:
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        return args.output
    name = src.with_suffix(".docx").name
    if args.output_dir:
        args.output_dir.mkdir(parents=True, exist_ok=True)
        return args.output_dir / name
    return src.with_suffix(".docx")


# ------------------------------------------------------------------- cli -----


class _HelpFormatter(
    argparse.RawDescriptionHelpFormatter,
    argparse.ArgumentDefaultsHelpFormatter,
):
    """Keep epilog formatting, show defaults next to every flag."""


_EPILOG = """\
examples:
  md2cv cv.md                                    # -> cv.docx next to it
  md2cv cv.md -o Oleksii_CV.docx                 # custom output path
  md2cv ~/cv/*.md -d ~/cv/docx                   # batch into a folder
  md2cv examples/                                # all *.md from a folder
  md2cv cv.md --preset compact                   # squeeze into two pages
  md2cv cv.md --preset comfortable               # one-pager with air
  md2cv cv.md --preset classic                   # sober Word-blue, no rules
  md2cv cv.md --preset compact --accent '#8B0000'  # flags override preset
  md2cv --list-presets                           # inspect preset contents

notes:
  · Explicit flags override --preset values, regardless of order.
  · --rule-size 0 disables the rule under section headings.
  · For PDFs: libreoffice --headless --convert-to pdf --outdir . cv.docx
"""


def build_parser() -> argparse.ArgumentParser:
    C = Config
    p = argparse.ArgumentParser(
        prog=PROG,
        description="Convert structured Markdown résumés to neatly typeset DOCX files.",
        epilog=_EPILOG,
        formatter_class=_HelpFormatter,
    )
    p.add_argument(
        "inputs",
        nargs="*",
        type=Path,
        metavar="FILE.md",
        help="Input files and/or directories (directories are scanned for *.md).",
    )
    p.add_argument(
        "-o",
        "--output",
        type=Path,
        help="Output path (only with a single input). Parent dirs are created.",
    )
    p.add_argument(
        "-d",
        "--output-dir",
        type=Path,
        help="Directory to write outputs into; created if missing.",
    )
    p.add_argument(
        "--preset",
        choices=sorted(PRESETS),
        default="default",
        help="Apply a bundle of layout defaults; explicit flags override it.",
    )
    p.add_argument(
        "--list-presets",
        action="store_true",
        help="Print available presets and exit.",
    )
    p.add_argument("--version", action="version", version=f"{PROG} {__version__}")

    g = p.add_argument_group("typography")
    g.add_argument("--font", default=C.font)
    g.add_argument("--body-size", type=float, default=C.body_size, metavar="PT")
    g.add_argument("--name-size", type=float, default=C.name_size, metavar="PT")
    g.add_argument("--tagline-size", type=float, default=C.tagline_size, metavar="PT")
    g.add_argument("--contact-size", type=float, default=C.contact_size, metavar="PT")
    g.add_argument("--section-size", type=float, default=C.section_size, metavar="PT")
    g.add_argument("--company-size", type=float, default=C.company_size, metavar="PT")
    g.add_argument("--role-size", type=float, default=C.role_size, metavar="PT")
    g.add_argument("--project-size", type=float, default=C.project_size, metavar="PT")
    g.add_argument("--line-spacing", type=float, default=C.line_spacing, metavar="F")

    g = p.add_argument_group("page geometry (cm)")
    g.add_argument("--page-width", type=float, default=C.page_width, metavar="CM")
    g.add_argument("--page-height", type=float, default=C.page_height, metavar="CM")
    g.add_argument("--margin-top", type=float, default=C.margin_top, metavar="CM")
    g.add_argument("--margin-bottom", type=float, default=C.margin_bottom, metavar="CM")
    g.add_argument("--margin-left", type=float, default=C.margin_left, metavar="CM")
    g.add_argument("--margin-right", type=float, default=C.margin_right, metavar="CM")

    g = p.add_argument_group("colors (hex, e.g. #1F3864)")
    g.add_argument("--accent", type=_hex, default=C.accent)
    g.add_argument("--muted", type=_hex, default=C.muted)
    g.add_argument("--rule-color", type=_hex, default=C.rule_color)
    g.add_argument(
        "--rule-size",
        type=int,
        default=C.rule_size,
        metavar="N",
        help="Rule thickness in eighths of a point; 0 disables it.",
    )

    g = p.add_argument_group("spacing (pt, except --list-indent in cm)")
    g.add_argument(
        "--section-space-before",
        type=float,
        default=C.section_space_before,
        metavar="PT",
    )
    g.add_argument(
        "--section-space-after",
        type=float,
        default=C.section_space_after,
        metavar="PT",
    )
    g.add_argument(
        "--company-space-before",
        type=float,
        default=C.company_space_before,
        metavar="PT",
    )
    g.add_argument(
        "--item-space-after",
        type=float,
        default=C.item_space_after,
        metavar="PT",
    )
    g.add_argument(
        "--para-space-after",
        type=float,
        default=C.para_space_after,
        metavar="PT",
    )
    g.add_argument("--list-indent", type=float, default=C.list_indent, metavar="CM")

    return p


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()

    pre, _ = parser.parse_known_args(argv)

    if pre.list_presets:
        for name, values in PRESETS.items():
            if values:
                print(f"{name}:")
                for k, v in values.items():
                    print(f"    {k} = {v}")
            else:
                print(f"{name}:  (built-in defaults, no overrides)")
            print()
        return 0

    if PRESETS[pre.preset]:
        parser.set_defaults(**PRESETS[pre.preset])

    args = parser.parse_args(argv)

    if not args.inputs:
        parser.error("no input files given (see --help for examples)")

    status = 0
    resolved: list[Path] = []
    for item in args.inputs:
        if item.is_dir():
            found = sorted(item.glob("*.md"))
            if not found:
                print(f"{PROG}: no *.md files in {item}", file=sys.stderr)
                status = 1
            resolved.extend(found)
        elif item.exists():
            resolved.append(item)
        else:
            print(f"{PROG}: not found: {item}", file=sys.stderr)
            status = 1

    if args.output and len(resolved) > 1:
        parser.error("-o/--output works only with a single input file")

    if not resolved:
        return status or 1

    cfg_keys = {f.name for f in fields(Config)}
    cfg = Config(**{k: v for k, v in vars(args).items() if k in cfg_keys})

    for src in resolved:
        dst = _resolve_output(src, args)
        try:
            convert(src, dst, cfg)
        except Exception as exc:
            print(f"{PROG}: failed to convert {src}: {exc}", file=sys.stderr)
            status = 1
            continue
        print(f"{PROG}: wrote {dst}")
    return status


if __name__ == "__main__":
    raise SystemExit(main())
