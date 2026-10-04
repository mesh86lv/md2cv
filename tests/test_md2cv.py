"""Regression tests for md2cv, focused on parser edge cases."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document as DocxDocument
from docx.enum.text import WD_TAB_ALIGNMENT

import md2cv

REPO_ROOT = Path(__file__).resolve().parent.parent
SAMPLE = REPO_ROOT / "examples" / "sample.md"


def _render(md_text: str, tmp_path: Path):
    src = tmp_path / "in.md"
    src.write_text(md_text, encoding="utf-8")
    out = tmp_path / "out.docx"
    md2cv.convert(src, out, md2cv.Config())
    return DocxDocument(str(out))


# -- fix 1: role line right after ### / ####, blank lines in between --------


def test_role_line_after_company_with_blank_line(tmp_path):
    doc = _render(
        "# N\n\n"
        "## Experience\n\n"
        "### Acme Corp\n\n"
        "**Senior Software Engineer** · Jan 2022 – Present\n\n"
        "- did a thing\n",
        tmp_path,
    )
    role_paragraphs = [p for p in doc.paragraphs if "\t" in p.text]
    assert role_paragraphs, "role line with tab was not produced"
    p = role_paragraphs[0]
    assert p.text.startswith("Senior Software Engineer")
    assert "Jan 2022" in p.text
    stops = list(p.paragraph_format.tab_stops)
    assert any(s.alignment == WD_TAB_ALIGNMENT.RIGHT for s in stops), (
        "role paragraph lacks a right-aligned tab stop"
    )


def test_role_line_after_project(tmp_path):
    doc = _render(
        "# N\n\n"
        "## Experience\n\n"
        "### Co\n\n"
        "**Engineer** · 2020 – 2021\n\n"
        "#### Project: Widget\n\n"
        "**Tech Lead** · 2021 – 2022\n",
        tmp_path,
    )
    role_paragraphs = [p for p in doc.paragraphs if "\t" in p.text]
    assert len(role_paragraphs) == 2
    assert "Tech Lead" in role_paragraphs[1].text


def test_sample_md_produces_role_lines(tmp_path):
    """The shipped example itself must round-trip through the role parser."""
    out = tmp_path / "sample.docx"
    md2cv.convert(SAMPLE, out, md2cv.Config())
    doc = DocxDocument(str(out))
    role_paragraphs = [p for p in doc.paragraphs if "\t" in p.text]
    assert len(role_paragraphs) >= 3, (
        f"expected at least 3 role lines, got {len(role_paragraphs)}"
    )


# -- fix 2: underscore italic ------------------------------------------------


def test_underscore_italic_is_parsed(tmp_path):
    doc = _render(
        "# N\n\n## S\n\n_Optional italic blurb._\n\n_Stack:_ Node.js, PostgreSQL\n",
        tmp_path,
    )
    italic_runs = [r.text for p in doc.paragraphs for r in p.runs if r.italic]
    assert "Optional italic blurb." in italic_runs
    assert "Stack:" in italic_runs


def test_snake_case_is_not_italicised(tmp_path):
    doc = _render("# N\n\n## S\n\nlook at my_var_name here\n", tmp_path)
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "my_var_name" in body
    italic_runs = [r.text for p in doc.paragraphs for r in p.runs if r.italic]
    assert not any("var" in r for r in italic_runs)


# -- fix 3: core properties --------------------------------------------------


def test_author_and_title_are_set(tmp_path):
    doc = _render("# Jane Doe\n\n## S\n\n- a\n", tmp_path)
    assert doc.core_properties.author == "Jane Doe"
    assert doc.core_properties.title == "Jane Doe"


# -- fix 4: BOM tolerance ----------------------------------------------------


def test_bom_does_not_break_name(tmp_path):
    src = tmp_path / "bom.md"
    src.write_bytes("\ufeff# Jane Doe\n\n## S\n\n- a\n".encode("utf-8"))
    out = tmp_path / "out.docx"
    md2cv.convert(src, out, md2cv.Config())
    doc = DocxDocument(str(out))
    assert doc.paragraphs[0].text == "Jane Doe"
    assert doc.core_properties.author == "Jane Doe"
